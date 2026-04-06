from fastapi import APIRouter, Depends, Query, Body, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text
from .web_database import get_db
from datetime import date, timedelta
from typing import Annotated, Optional, List
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from .bot.engine import bot
from aiogram.types import BufferedInputFile
from utils.exchange_rates import get_exchange_rates, convert_currency
from database.db_query import db as app_db

router = APIRouter(prefix="/stats", tags=["Stats"])

@router.get("/settings/{telegram_id}")
async def get_user_settings(telegram_id: int, db: AsyncSession = Depends(get_db)):
    q = text("SELECT fullname, sex, social_status, language, currency FROM users WHERE telegram_id = :uid")
    r = (await db.execute(q, {"uid": telegram_id})).mappings().first()
    if not r:
        return {"fullname": "", "sex": "male", "social_status": "other", "language": "uz", "currency": "uzs"}
    return dict(r)

@router.post("/settings/{telegram_id}")
async def update_user_settings(
    telegram_id: int,
    data: dict = Body(...),
    db: AsyncSession = Depends(get_db)
):
    q = text("""
        UPDATE users
        SET fullname = :fullname, sex = :sex, social_status = :social_status,
            language = :language, currency = :currency
        WHERE telegram_id = :uid
    """)
    await db.execute(q, {
        "uid": telegram_id,
        "fullname": data.get("fullname"),
        "sex": data.get("sex"),
        "social_status": data.get("social_status"),
        "language": data.get("language"),
        "currency": data.get("currency")
    })
    await db.commit()
    return {"status": "success"}

@router.get("/categories/{telegram_id}")
async def get_categories(telegram_id: int, db: AsyncSession = Depends(get_db)):
    # We use the app_db helper to handle initialization if needed
    rows = await app_db.get_categories(telegram_id, 'income')
    rows += await app_db.get_categories(telegram_id, 'expense')

    # Refresh from DB to get IDs and all fields
    q = text("SELECT id, type, name, emoji FROM custom_categories WHERE user_id = :uid")
    rows = (await db.execute(q, {"uid": telegram_id})).mappings().all()
    return [dict(r) for r in rows]

@router.post("/categories/{telegram_id}")
async def add_category(
    telegram_id: int,
    data: dict = Body(...),
    db: AsyncSession = Depends(get_db)
):
    # Check limit
    count_q = text("SELECT COUNT(*) FROM custom_categories WHERE user_id = :uid AND type = :type")
    count = await db.execute(count_q, {"uid": telegram_id, "type": data.get("type")})
    if count.scalar() >= 15:
        raise HTTPException(status_code=400, detail="Maximum 15 categories allowed per type")

    q = text("INSERT INTO custom_categories(user_id, type, name, emoji) VALUES(:uid, :type, :name, :emoji)")
    await db.execute(q, {
        "uid": telegram_id,
        "type": data.get("type"),
        "name": data.get("name"),
        "emoji": data.get("emoji")
    })
    await db.commit()
    return {"status": "success"}

@router.put("/categories/{telegram_id}/{cat_id}")
async def update_category(
    telegram_id: int,
    cat_id: int,
    data: dict = Body(...),
    db: AsyncSession = Depends(get_db)
):
    q = text("UPDATE custom_categories SET name = :name, emoji = :emoji WHERE id = :cid AND user_id = :uid")
    await db.execute(q, {"cid": cat_id, "uid": telegram_id, "name": data.get("name"), "emoji": data.get("emoji")})
    await db.commit()
    return {"status": "success"}

@router.delete("/categories/{telegram_id}/{cat_id}")
async def delete_category(
    telegram_id: int,
    cat_id: int,
    db: AsyncSession = Depends(get_db)
):
    # Check min 3
    # First get type of this category
    type_q = text("SELECT type FROM custom_categories WHERE id = :cid")
    cat_type = (await db.execute(type_q, {"cid": cat_id})).scalar()

    count_q = text("SELECT COUNT(*) FROM custom_categories WHERE user_id = :uid AND type = :type")
    count = await db.execute(count_q, {"uid": telegram_id, "type": cat_type})
    if count.scalar() <= 3:
        raise HTTPException(status_code=400, detail="At least 3 categories must remain")

    q = text("DELETE FROM custom_categories WHERE id = :cid AND user_id = :uid")
    await db.execute(q, {"cid": cat_id, "uid": telegram_id})
    await db.commit()
    return {"status": "success"}

async def get_user_currency(telegram_id: int, db: AsyncSession) -> str:
    q = text("SELECT currency FROM users WHERE telegram_id = :uid")
    res = await db.execute(q, {"uid": telegram_id})
    val = res.scalar()
    return val or "usd"

@router.get("/summary/{telegram_id}")
async def monthly_summary(
    telegram_id: int,
    period: str = Query("month", pattern="^(day|month)$"),
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)

    # We need to fetch sums grouped by currency
    async def get_converted_sum(table, payment_method=None):
        if period == "day":
            where = "user_id=:uid AND inserted_at::date = now()::date"
        else:
            where = "user_id=:uid AND inserted_at >= date_trunc('month', now())"

        if payment_method:
            where += f" AND payment_method='{payment_method}'"

        q = text(f"SELECT currency, SUM(amount) as total FROM {table} WHERE {where} GROUP BY currency")
        rows = (await db.execute(q, {"uid": telegram_id})).mappings().all()

        grand_total = 0.0
        for row in rows:
            amt = float(row['total'])
            curr = row['currency']
            converted = await convert_currency(amt, curr, target_curr)
            grand_total += converted
        return grand_total

    income = await get_converted_sum("user_earnings")
    income_cash = await get_converted_sum("user_earnings", "cash")
    income_card = await get_converted_sum("user_earnings", "card")

    expense = await get_converted_sum("user_expenses")
    expense_cash = await get_converted_sum("user_expenses", "cash")
    expense_card = await get_converted_sum("user_expenses", "card")

    return {
        "income": income,
        "income_cash": income_cash,
        "income_card": income_card,
        "expense": expense,
        "expense_cash": expense_cash,
        "expense_card": expense_card,
        "result": income - expense,
        "currency": target_curr.upper()
    }

@router.get("/trend/{telegram_id}")
async def trend(telegram_id: int, db: AsyncSession = Depends(get_db)):
    target_curr = await get_user_currency(telegram_id, db)

    # Generate series of months
    q_months = text("""
        SELECT to_char(m,'Mon') AS month_name, CAST(date_trunc('month', m) AS TIMESTAMP) as month_date
        FROM generate_series(
            date_trunc('month', now()) - interval '11 months',
            date_trunc('month', now()),
            interval '1 month'
        ) m
        ORDER BY m
    """)
    months = (await db.execute(q_months)).mappings().all()

    result = []
    for m in months:
        # For each month, get income and expense sums grouped by currency
        async def get_monthly_sum(table, month_start):
            q = text(f"""
                SELECT currency, SUM(amount) as total
                FROM {table}
                WHERE user_id = :uid AND date_trunc('month', inserted_at) = :m
                GROUP BY currency
            """)
            rows = (await db.execute(q, {"uid": telegram_id, "m": month_start})).mappings().all()
            total = 0.0
            for r in rows:
                total += await convert_currency(float(r['total']), r['currency'], target_curr)
            return total

        m_date = m['month_date']
        if hasattr(m_date, 'replace'):
            m_date = m_date.replace(tzinfo=None)
        income = await get_monthly_sum("user_earnings", m_date)
        expense = await get_monthly_sum("user_expenses", m_date)

        result.append({
            "month": m['month_name'],
            "income": income,
            "expense": expense
        })
    return result

@router.get("/expenses-pie/{telegram_id}")
async def expenses_pie(
        telegram_id: int,
        date_from: date,
        date_to: date,
        db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    dt_plus_one = date_to + timedelta(days=1)
    q = text("""
        SELECT source, currency, SUM(amount) total
        FROM user_expenses
        WHERE user_id = :uid
          AND inserted_at >= :df
          AND inserted_at < :dt
        GROUP BY source, currency
    """)
    rows = (await db.execute(q, {
        "uid": telegram_id,
        "df": date_from,
        "dt": dt_plus_one
    })).mappings().all()

    cat_totals = {}
    for r in rows:
        converted = await convert_currency(float(r['total']), r['currency'], target_curr)
        cat_totals[r['source']] = cat_totals.get(r['source'], 0.0) + converted

    res = [{"source": k, "total": v} for k, v in cat_totals.items()]
    res.sort(key=lambda x: x['total'], reverse=True)
    return res

@router.get("/monitoring/{telegram_id}")
async def monitoring(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"

    async def get_today_sum(payment_method=None):
        where = "user_id = :uid AND inserted_at::date = now()::date"
        if payment_method:
            where += f" AND payment_method = '{payment_method}'"
        q = text(f"SELECT currency, SUM(amount) as total FROM {table} WHERE {where} GROUP BY currency")
        rows = (await db.execute(q, {"uid": telegram_id})).mappings().all()
        total = 0.0
        for r in rows:
            total += await convert_currency(float(r['total']), r['currency'], target_curr)
        return total

    total = await get_today_sum()
    cash = await get_today_sum("cash")
    card = await get_today_sum("card")

    q_cat = text(f"""
        SELECT source, currency, SUM(amount) as total
        FROM {table}
        WHERE user_id = :uid AND inserted_at::date = now()::date
        GROUP BY source, currency
    """)
    rows = (await db.execute(q_cat, {"uid": telegram_id})).mappings().all()

    cat_totals = {}
    for r in rows:
        src = r['source']
        converted = await convert_currency(float(r['total']), r['currency'], target_curr)
        cat_totals[src] = cat_totals.get(src, 0.0) + converted

    categories = [{"source": k, "total": v} for k, v in cat_totals.items()]
    categories.sort(key=lambda x: x['total'], reverse=True)

    q_trans = text(f"""
        SELECT id, amount, currency, source, payment_method, additional_info, to_char(inserted_at, 'HH24:MI') as time
        FROM {table}
        WHERE user_id = :uid AND inserted_at::date = now()::date
        ORDER BY inserted_at DESC
    """)
    trans_rows = (await db.execute(q_trans, {"uid": telegram_id})).mappings().all()
    transactions = []
    for r in trans_rows:
        d = dict(r)
        d['amount'] = await convert_currency(float(r['amount']), r['currency'], target_curr)
        transactions.append(d)

    return {
        "summary": {"total": total, "cash": cash, "card": card},
        "categories": categories,
        "transactions": transactions
    }

@router.get("/report/{telegram_id}")
async def report(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    date_from: date,
    date_to: date,
    source: Optional[str] = None,
    payment_method: Optional[str] = None,
    sort_by: str = "inserted_at",
    order: str = "DESC",
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"
    dt_plus_one = date_to + timedelta(days=1)

    conditions = ["user_id = :uid", "inserted_at >= :df", "inserted_at < :dt"]
    params = {"uid": telegram_id, "df": date_from, "dt": dt_plus_one}

    if source:
        conditions.append("source = :source")
        params["source"] = source
    if payment_method:
        conditions.append("payment_method = :pm")
        params["pm"] = payment_method

    where_clause = " AND ".join(conditions)

    if sort_by not in ["inserted_at", "amount", "source", "payment_method"]:
        sort_by = "inserted_at"
    if order not in ["ASC", "DESC"]:
        order = "DESC"

    q = text(f"""
        SELECT id, amount, currency, source, payment_method, additional_info, to_char(inserted_at, 'YYYY-MM-DD HH24:MI') as date
        FROM {table}
        WHERE {where_clause}
        ORDER BY {sort_by} {order}
    """)

    rows = (await db.execute(q, params)).mappings().all()

    result = []
    for r in rows:
        d = dict(r)
        d['amount'] = await convert_currency(float(r['amount']), r['currency'], target_curr)
        d['original_amount'] = float(r['amount'])
        d['original_currency'] = r['currency']
        result.append(d)

    # Re-sort if sorted by amount after conversion
    if sort_by == "amount":
        result.sort(key=lambda x: x['amount'], reverse=(order == "DESC"))

    return result

@router.post("/transaction/{telegram_id}/{type}")
async def add_web_transaction(
    telegram_id: int,
    type: str,
    data: dict = Body(...),
    db: AsyncSession = Depends(get_db)
):
    table = "user_earnings" if type == "income" else "user_expenses"
    q = text(f"""
        INSERT INTO {table}(user_id, amount, currency, source, payment_method, additional_info)
        VALUES(:uid, :amount, :currency, :source, :pm, :info)
    """)
    await db.execute(q, {
        "uid": telegram_id,
        "amount": float(data.get("amount")),
        "currency": data.get("currency", "uzs"),
        "source": data.get("source"),
        "pm": data.get("payment_method", "cash"),
        "info": data.get("additional_info")
    })
    await db.commit()
    return {"status": "success"}

@router.delete("/transaction/{telegram_id}/{type}/{transaction_id}")
async def delete_transaction(
    telegram_id: int,
    type: str,
    transaction_id: int,
    db: AsyncSession = Depends(get_db)
):
    table = "user_earnings" if type == "income" else "user_expenses"
    q = text(f"DELETE FROM {table} WHERE id = :tid AND user_id = :uid")
    await db.execute(q, {"tid": transaction_id, "uid": telegram_id})
    await db.commit()
    return {"status": "success"}

@router.put("/transaction/{telegram_id}/{type}/{transaction_id}")
async def update_transaction(
    telegram_id: int,
    type: str,
    transaction_id: int,
    data: dict = Body(...),
    db: AsyncSession = Depends(get_db)
):
    table = "user_earnings" if type == "income" else "user_expenses"
    q = text(f"""
        UPDATE {table}
        SET amount = :amount, source = :source, payment_method = :pm, additional_info = :info
        WHERE id = :tid AND user_id = :uid
    """)
    await db.execute(q, {
        "amount": float(data.get("amount")),
        "source": data.get("source"),
        "pm": data.get("payment_method"),
        "info": data.get("additional_info"),
        "tid": transaction_id,
        "uid": telegram_id
    })
    await db.commit()
    return {"status": "success"}

@router.get("/export/{telegram_id}")
async def export_excel(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    date_from: date,
    date_to: date,
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"
    dt_plus_one = date_to + timedelta(days=1)

    q = text(f"""
        SELECT to_char(inserted_at, 'YYYY-MM-DD HH24:MI') as date, amount, currency, source, payment_method, additional_info
        FROM {table}
        WHERE user_id = :uid AND inserted_at >= :df AND inserted_at < :dt
        ORDER BY inserted_at DESC
    """)
    rows = (await db.execute(q, {"uid": telegram_id, "df": date_from, "dt": dt_plus_one})).mappings().all()

    if not rows:
        raise HTTPException(status_code=404, detail="No data found for the selected period")

    # Create Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"

    headers = ["Sana", "Summa", "Kategoriya", "To'lov usuli", "Izoh"]
    ws.append(headers)

    # Style headers
    header_fill = PatternFill(start_color="3498DB", end_color="3498DB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        amt = await convert_currency(float(row['amount']), row['currency'], target_curr)
        ws.append([row['date'], amt, row['source'], row['payment_method'], row['additional_info']])

    # Auto-adjust columns width
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    # Save to buffer
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    # Send via Bot
    filename = f"hisobot_{type}_{date_from}_{date_to}.xlsx"
    input_file = BufferedInputFile(excel_file.read(), filename=filename)

    try:
        await bot.send_document(chat_id=telegram_id, document=input_file, caption=f"📊 Sizning {date_from} dan {date_to} gacha bo'lgan {type} hisobotingiz.")
        return {"status": "success", "message": "Excel file sent to Telegram"}
    except Exception as e:
        print(f"Error sending document: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to send document: {str(e)}")

@router.get("/weekly/{telegram_id}")
async def weekly_stats(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    date_from: date,
    date_to: date,
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"
    dt_plus_one = date_to + timedelta(days=1)

    q = text(f"""
        SELECT EXTRACT(DOW FROM inserted_at) as dow, currency, SUM(amount) as total
        FROM {table}
        WHERE user_id = :uid AND inserted_at >= :df AND inserted_at < :dt
        GROUP BY dow, currency
    """)
    rows = (await db.execute(q, {"uid": telegram_id, "df": date_from, "dt": dt_plus_one})).mappings().all()

    day_totals = {float(i): 0.0 for i in range(7)}
    for r in rows:
        converted = await convert_currency(float(r['total']), r['currency'], target_curr)
        day_totals[float(r['dow'])] += converted

    days_map = {0: 'Yak', 1: 'Dush', 2: 'Sesh', 3: 'Chor', 4: 'Pay', 5: 'Jum', 6: 'Shan'}
    res = []
    # Order: Mon (1) to Sun (0)
    for i in [1, 2, 3, 4, 5, 6, 0]:
        res.append({"day": days_map[i], "total": day_totals[float(i)]})
    return res

@router.get("/daily/{telegram_id}")
async def daily_stats(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    date_from: date,
    date_to: date,
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"

    q = text(f"""
        SELECT inserted_at::date as day, currency, SUM(amount) as total
        FROM {table}
        WHERE user_id = :uid AND inserted_at::date >= :df AND inserted_at::date <= :dt
        GROUP BY day, currency
    """)
    rows = (await db.execute(q, {"uid": telegram_id, "df": date_from, "dt": date_to})).mappings().all()

    daily_map = {}
    for r in rows:
        day_str = r['day'].isoformat()
        converted = await convert_currency(float(r['total']), r['currency'], target_curr)
        daily_map[day_str] = daily_map.get(day_str, 0.0) + converted

    res = []
    curr = date_from
    while curr <= date_to:
        d_str = curr.isoformat()
        res.append({"day": d_str, "total": daily_map.get(d_str, 0.0)})
        curr += timedelta(days=1)
    return res

@router.get("/metrics/{telegram_id}")
async def metrics(
    telegram_id: int,
    type: Annotated[str, Query(pattern="^(income|expenses)$")],
    date_from: date,
    date_to: date,
    db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    table = "user_earnings" if type == "income" else "user_expenses"
    dt_plus_one = date_to + timedelta(days=1)

    q = text(f"SELECT amount, currency FROM {table} WHERE user_id = :uid AND inserted_at >= :df AND inserted_at < :dt")
    rows = (await db.execute(q, {"uid": telegram_id, "df": date_from, "dt": dt_plus_one})).mappings().all()

    if not rows:
        return {"average": 0.0, "count": 0, "max": 0.0, "total": 0.0}

    converted_amounts = []
    for r in rows:
        converted_amounts.append(await convert_currency(float(r['amount']), r['currency'], target_curr))

    return {
        "average": sum(converted_amounts) / len(converted_amounts),
        "count": len(converted_amounts),
        "max": max(converted_amounts),
        "total": sum(converted_amounts)
    }

@router.get("/income-pie/{telegram_id}")
async def income_pie(
        telegram_id: int,
        date_from: date,
        date_to: date,
        db: AsyncSession = Depends(get_db)
):
    target_curr = await get_user_currency(telegram_id, db)
    dt_plus_one = date_to + timedelta(days=1)
    q = text("""
        SELECT source, currency, SUM(amount) total
        FROM user_earnings
        WHERE user_id = :uid
          AND inserted_at >= :df
          AND inserted_at < :dt
        GROUP BY source, currency
    """)
    rows = (await db.execute(q, {
        "uid": telegram_id,
        "df": date_from,
        "dt": dt_plus_one
    })).mappings().all()

    cat_totals = {}
    for r in rows:
        converted = await convert_currency(float(r['total']), r['currency'], target_curr)
        cat_totals[r['source']] = cat_totals.get(r['source'], 0.0) + converted

    res = [{"source": k, "total": v} for k, v in cat_totals.items()]
    res.sort(key=lambda x: x['total'], reverse=True)
    return res
